# 离线行为测试：真实临时 SQLite/Git + 替换模型和构建结果，不改正式 CRM。
# 阅读这些测试可理解 MVP 的边界：幂等、权限、修复上限、中断和串行执行。
import subprocess
import threading
from io import BytesIO
from types import SimpleNamespace

import pytest
from app.agent import Budget
from app.config import Settings
from app.documents import parse_prd, validate_image
from app.main import create_app
from app.repository import Repository
from app.store import Store
from app.worker import Worker
from app.workflow import Workflow
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image


def image_bytes():
    # 生成最小有效 PNG，让上传校验走真实 Pillow，不依赖仓库里的设计图。
    buf = BytesIO()
    Image.new("RGB", (20, 20), "white").save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def settings(tmp_path):
    # pytest 为每个测试创建独立临时目录，构造 main 基线和假依赖目录。
    # Settings 注入假密钥只是通过配置检查，不代表会调用任何外部模型。
    target = tmp_path / "crm"
    (target / "src").mkdir(parents=True)
    (target / "node_modules").mkdir()
    (target / "README.md").write_text("React demo rules")
    (target / "src/App.tsx").write_text(
        "export default function App(){return <p>Baseline</p>}"
    )

    def git(*args):
        return subprocess.run(
            ["git", *args], cwd=target, check=True, capture_output=True, text=True
        ).stdout.strip()

    git("init", "-b", "main")
    git("add", "README.md", "src")
    git(
        "-c",
        "user.name=Test",
        "-c",
        "user.email=test@localhost",
        "commit",
        "-m",
        "baseline",
    )
    return Settings(
        root=tmp_path,
        target=target,
        runtime=tmp_path / "runtime",
        api_key="fake-key-tests-only",
    )


def create_task(settings, store, task_id="t-111111111111"):
    # 模拟“HTTP 已保存材料并写入队列”的起点，供工作流测试直接使用。
    directory = settings.runtime / "tasks" / task_id
    directory.mkdir(parents=True)
    (directory / "design.png").write_bytes(image_bytes())
    (directory / "prd").write_text("给客户列表新增筛选。")
    return store.create(
        task_id,
        "key-" + task_id,
        "fingerprint",
        {
            "prd_name": "prd.md",
            "design_name": "design.png",
            "design_suffix": ".png",
            "instructions": "",
            "checks": [],
        },
    )


def test_file_validation_and_parsers():
    # 验证真实图片校验、拒绝无效类型，以及 DOCX 段落/表格文字提取。
    validate_image(image_bytes(), "design.png")
    with pytest.raises(ValueError):
        validate_image(b"not an image", "design.png")
    with pytest.raises(ValueError):
        validate_image(image_bytes(), "design.svg")
    assert parse_prd("需求".encode(), "prd.md") == "需求"
    with pytest.raises(ValueError):
        parse_prd(b"", "prd.md")
    with pytest.raises(ValueError):
        parse_prd(b"hello", "prd.exe")
    document = Document()
    document.add_paragraph("新增状态筛选")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "状态"
    table.cell(0, 1).text = "已成交"
    buffer = BytesIO()
    document.save(buffer)
    text = parse_prd(buffer.getvalue(), "prd.docx")
    assert "新增状态筛选" in text and "已成交" in text


def test_upload_idempotency_and_validation(settings):
    # TestClient 走完整 API，但关闭 Worker，因此重复 POST 不会触发模型执行。
    # 同 key 同材料应只有一个任务；改材料返回 409；任意 .env 下载应被拒绝。
    app = create_app(settings, start_worker=False)
    with TestClient(app) as client:
        files = {
            "design": ("design.png", image_bytes(), "image/png"),
            "prd": ("prd.md", "需求".encode(), "text/markdown"),
        }
        body = {"idempotency_key": "same-key-123", "instructions": ""}
        assert client.post("/api/tasks", files=files, data=body).status_code == 403
        headers = {"X-Design-Agent": "local-mvp"}
        a = client.post("/api/tasks", headers=headers, files=files, data=body)
        b = client.post("/api/tasks", headers=headers, files=files, data=body)
        assert (
            a.status_code == b.status_code == 202 and a.json()["id"] == b.json()["id"]
        )
        assert len(app.state.store.list()) == 1
        assert (
            client.get("/api/submissions/same-key-123").json()["id"] == a.json()["id"]
        )
        altered = {**body, "instructions": "different"}
        assert (
            client.post(
                "/api/tasks", headers=headers, files=files, data=altered
            ).status_code
            == 409
        )
        bad = {**files, "design": ("design.png", b"invalid", "image/png")}
        assert (
            client.post(
                "/api/tasks",
                headers=headers,
                files=bad,
                data={"idempotency_key": "another-key"},
            ).status_code
            == 422
        )
        assert client.get("/api/tasks/t-111111111111").status_code == 404
        assert (
            client.get(f"/api/tasks/{a.json()['id']}/artifacts/.env").status_code == 404
        )


def test_repository_isolation_and_paths(settings):
    # 真正创建 Git worktree，验证越权路径/符号链接被拒绝、新文件能进入 diff。
    # 同时确认主目录和 main 提交保持原基线，避免只测“工具返回成功”。
    repo = Repository(
        settings.target, settings.runtime / "worktrees" / "t-222222222222"
    )
    before = Repository.command(["git", "rev-parse", "main"], settings.target)
    repo.prepare("t-222222222222")
    assert "node_modules" not in Repository.command(
        ["git", "status", "--porcelain"], repo.worktree
    )
    for path in ["../secrets", ".env", "package.json", "/tmp/no", "src/../../secret"]:
        with pytest.raises(ValueError):
            repo.path(path, write=True)
    (repo.worktree / "src/escape.ts").symlink_to(settings.target / "README.md")
    with pytest.raises(ValueError):
        repo.write("src/escape.ts", "no")
    (repo.worktree / "src/escape.ts").unlink()
    repo.write("src/New.tsx", 'export const label="Feature";')
    diff, files = repo.diff()
    assert "src/New.tsx" in files and "Feature" in diff
    assert not (settings.target / "src/New.tsx").exists()
    assert Repository.command(["git", "rev-parse", "main"], settings.target) == before
    assert (
        Repository.command(
            ["git", "status", "--porcelain", "--untracked-files=no"], settings.target
        )
        == ""
    )


class FakeDeveloper:
    # 保留与 Developer 一样的接口，确定性写一处源码，去掉模型随机性和费用。
    # 这里的 budget 是计数替身，不能作为真实模型 token 使用证据。
    def __init__(self, settings, repo, event):
        self.repo = repo
        self.budget = SimpleNamespace(calls=0, tokens=0)

    def plan(self, prd, design):
        self.budget.calls += 1
        return {
            "summary": "Offline test fixture",
            "files": ["src/App.tsx"],
            "steps": ["edit"],
            "acceptance": ["fixture"],
        }

    def develop(self, *args):
        self.budget.calls += 1
        self.repo.write(
            "src/App.tsx",
            "export default function App(){return <p>Changed in offline test</p>}",
        )
        return "Offline fixture only, no real model call."


async def fake_vision(*args):
    # 注入异步替身，仍验证 materials 用 asyncio.run 等待结果的调用链。
    return "Offline visual fixture only."


def test_graph_repairs_and_delivery(settings, monkeypatch):
    # 人为让首项检查失败、下一轮通过，验证只修复一次并成功保存真实 Git diff。
    store = Store(settings.runtime / "tasks.sqlite3")
    create_task(settings, store)
    counts = {"n": 0}

    def check(self, name):
        counts["n"] += 1
        return {
            "name": name,
            "status": "failed" if counts["n"] == 1 else "passed",
            "output": "offline check fixture",
        }

    monkeypatch.setattr(Repository, "check", check)
    workflow = Workflow(
        settings,
        store,
        "t-111111111111",
        vision=fake_vision,
        developer_factory=FakeDeveloper,
    )
    workflow.run()
    task = store.get("t-111111111111")
    assert task["status"] == "succeeded" and task["report"]["repair_rounds"] == 1
    assert task["report"]["model_calls"] == 3
    assert task["checks"][-1]["status"] == "not_run"
    assert "Baseline" in (settings.target / "src/App.tsx").read_text()
    assert (settings.runtime / "tasks/t-111111111111/changes.diff").is_file()


def test_graph_stops_on_check_failures(settings, monkeypatch):
    # 每轮都失败时必须达到上限后结束，保留改动与失败报告，不能无限循环。
    store = Store(settings.runtime / "tasks.sqlite3")
    create_task(settings, store)
    monkeypatch.setattr(
        Repository,
        "check",
        lambda self, name: {
            "name": name,
            "status": "failed",
            "output": "fixture fails",
        },
    )
    Workflow(
        settings,
        store,
        "t-111111111111",
        vision=fake_vision,
        developer_factory=FakeDeveloper,
    ).run()
    task = store.get("t-111111111111")
    assert task["status"] == "failed" and task["report"]["repair_rounds"] == 1
    assert task["report"]["changed_files"] == ["src/App.tsx"]


def test_vision_failure_and_restart_are_honest(settings):
    # 视觉异常与服务重启都应形成真实 failed/interrupted 状态，而不是虚报成功。
    store = Store(settings.runtime / "tasks.sqlite3")
    create_task(settings, store)

    async def failure(*args):
        raise RuntimeError("vision denied")

    Workflow(
        settings,
        store,
        "t-111111111111",
        vision=failure,
        developer_factory=FakeDeveloper,
    ).run()
    assert store.get("t-111111111111")["status"] == "failed"
    assert "vision denied" in store.get("t-111111111111")["error"]
    store.update("t-111111111111", status="running", branch="preserved")
    store.interrupt_running()
    task = store.get("t-111111111111")
    assert task["phase"] == "interrupted" and task["branch"] == "preserved"


def test_call_budget_and_redaction(settings):
    # 达到调用上限后下一次 consume 必须拒绝；日志中的当前配置密钥应被替换。
    settings.max_model_calls = 2
    budget = Budget(settings)
    budget.consume()
    budget.consume()
    with pytest.raises(RuntimeError):
        budget.consume()
    assert settings.api_key not in settings.redact("failure " + settings.api_key)


def test_worker_runs_without_browser_and_serializes_tasks(settings, monkeypatch):
    # Event 精确控制第一条任务何时结束，验证第二条不能在第一条阻塞时开始。
    # 不依赖浏览器轮询推动执行，证明 Worker 与网页请求生命周期独立。
    store = Store(settings.runtime / "tasks.sqlite3")
    create_task(settings, store)
    create_task(settings, store, "t-333333333333")
    first_started, release, finished = (
        threading.Event(),
        threading.Event(),
        threading.Event(),
    )
    starts = []

    class ControlledWorkflow:
        def __init__(self, config, database, task_id):
            self.store, self.task_id = database, task_id

        def run(self):
            starts.append(self.task_id)
            self.store.update(self.task_id, status="running", phase="develop")
            first_started.set()
            assert release.wait(2)
            self.store.update(self.task_id, status="succeeded", phase="completed")
            if len(starts) == 2:
                finished.set()

    monkeypatch.setattr("app.worker.Workflow", ControlledWorkflow)
    worker = Worker(settings, store)
    try:
        worker.start()
        assert first_started.wait(2)
        assert len(starts) == 1
        assert store.get("t-333333333333")["status"] == "queued"
        # 两条任务都由 Worker 自己推进，没有浏览器/轮询请求参与。
        release.set()
        assert finished.wait(2)
        assert all(task["status"] == "succeeded" for task in store.list())
    finally:
        release.set()
        worker.stop()
        worker.thread.join(timeout=2)
